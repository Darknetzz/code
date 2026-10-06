#!/usr/bin/env python3
"""
pyredact.py

Redact sensitive information from Scandinavian financial exports.

- CSV/TSV: keeps a whitelist of common columns; scrubs sensitive values in-place.
- TXT/OFX/QFX/QIF: scrubs the whole file in-place (plain-text bank dumps).
- JSON: walks objects/arrays and scrubs string (and sensitive numeric) values.
- XML: scrubs element text, tails, and attributes.
- HTML/HTM: scrubs text nodes and attributes (BeautifulSoup).
- XLSX: scrubs cell values across all sheets (openpyxl).
- PDF: uses PyMuPDF redaction annotations so matched text is removed from the
  content stream (not merely covered by black boxes).
- Images (PNG/JPEG/TIFF/BMP/GIF/WebP) and scanned PDF pages are OCR'd with
  Tesseract (tessdata) and the matched pixels are blanked.
- Scanned/image-only PDFs are skipped only when OCR is unavailable — they are
  never written to the output folder in that case.
- Redacts Norwegian personal numbers, account numbers, payment references, and
  any 9+ digit number. Optional templates cover email, phone, and URLs.
  Extra rules: --template / --pattern, or an interactive checklist.

Requires:
    pip install rich pymupdf typer openpyxl beautifulsoup4 questionary

Usage:
    python pyredact.py --help
    python pyredact.py
    python pyredact.py -i ./exports -o ./redacted
    python pyredact.py -i ./exports -r --dry-run
    python pyredact.py -i ./exports -t email -t phone -t url
    python pyredact.py -i ./exports --replace-patterns -t email -p "\\bIBAN[:\\s]*[A-Z0-9]+=>[IBAN]"
    python pyredact.py -i statement.pdf --dry-run
"""

from __future__ import annotations

import csv
import json
import re
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Annotated, Any

import typer
from rich.console import Console
from rich.panel import Panel
from rich.progress import BarColumn, Progress, SpinnerColumn, TaskProgressColumn, TextColumn
from rich.prompt import Confirm, Prompt
from rich.table import Table

app = typer.Typer(
    add_completion=False,
    rich_markup_mode="rich",
    help=(
        "Redact sensitive fields from Scandinavian financial exports "
        "(CSV/TSV/TXT/JSON/XML/HTML/XLSX/PDF/images and OFX/QIF).\n\n"
        "Missing [bold]--input[/bold] / [bold]--output[/bold] are prompted interactively. "
        "Pattern templates are a checklist unless [bold]--template[/bold] is given."
    ),
)

console = Console()

CSV_SUFFIXES = {".csv", ".tsv"}
TEXT_SUFFIXES = {".txt", ".ofx", ".qfx", ".qif"}
JSON_SUFFIXES = {".json"}
XML_SUFFIXES = {".xml"}
HTML_SUFFIXES = {".html", ".htm"}
XLSX_SUFFIXES = {".xlsx"}
PDF_SUFFIXES = {".pdf"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
SUPPORTED_SUFFIXES = (
    CSV_SUFFIXES
    | TEXT_SUFFIXES
    | JSON_SUFFIXES
    | XML_SUFFIXES
    | HTML_SUFFIXES
    | XLSX_SUFFIXES
    | PDF_SUFFIXES
    | IMAGE_SUFFIXES
)
OCR_DPI = 300
DEFAULT_OCR_LANG = "eng+nor"
_FILE_COUNT_ORDER = (
    "CSV",
    "TSV",
    "TXT",
    "OFX",
    "QFX",
    "QIF",
    "JSON",
    "XML",
    "HTML",
    "HTM",
    "XLSX",
    "PDF",
    "PNG",
    "JPG",
    "JPEG",
    "TIF",
    "TIFF",
    "BMP",
    "GIF",
    "WEBP",
)
_FILE_COUNT_STYLE = {
    "CSV": "cyan",
    "TSV": "cyan",
    "TXT": "green",
    "OFX": "green",
    "QFX": "green",
    "QIF": "green",
    "JSON": "blue",
    "XML": "blue",
    "HTML": "blue",
    "HTM": "blue",
    "XLSX": "yellow",
    "PDF": "magenta",
    "PNG": "bright_magenta",
    "JPG": "bright_magenta",
    "JPEG": "bright_magenta",
    "TIF": "bright_magenta",
    "TIFF": "bright_magenta",
    "BMP": "bright_magenta",
    "GIF": "bright_magenta",
    "WEBP": "bright_magenta",
}
# Pages with fewer alphanumeric chars than this are treated as having no usable text layer.
MIN_PAGE_ALNUM = 20

KEEP = {
    "dato",
    "date",
    "tekst",
    "beskrivelse",
    "description",
    "beløp",
    "belop",
    "amount",
    "inn",
    "ut",
}

PATTERN_SEP = "=>"


@dataclass(frozen=True)
class PatternTemplate:
    id: str
    description: str
    regex: str
    replacement: str
    flags: int = 0


TEMPLATES: tuple[PatternTemplate, ...] = (
    PatternTemplate(
        "id",
        "Norwegian personal number (fødselsnummer)",
        r"\b\d{6}\s?\d{5}\b",
        "[ID]",
    ),
    PatternTemplate(
        "acct",
        "Norwegian account number (kontonr)",
        r"\b\d{4}[. ]?\d{2}[. ]?\d{5}\b",
        "[ACCT]",
    ),
    PatternTemplate(
        "kid",
        "Payment reference (KID)",
        r"\bKID[:\s]*\d+",
        "KID",
        flags=re.I,
    ),
    PatternTemplate(
        "num",
        "Any other 9+ digit number",
        r"\b\d{9,}\b",
        "[NUM]",
    ),
    PatternTemplate(
        "email",
        "Email address",
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        "[EMAIL]",
    ),
    PatternTemplate(
        "phone",
        "Phone number (Nordic +47/+46/+45/+358, or other +country)",
        r"(?<!\d)\+(?:47|46|45|358|[1-9]\d{0,2})(?:[\s./-]*\d){6,12}(?!\d)",
        "[PHONE]",
    ),
    PatternTemplate(
        "url",
        "Webpage URL (http/https or www)",
        r"(?:https?://|www\.)[^\s<>\"']+",
        "[URL]",
        flags=re.I,
    ),
)
TEMPLATES_BY_ID = {t.id: t for t in TEMPLATES}
FINANCE_TEMPLATE_IDS = ("id", "acct", "kid", "num")
TEMPLATE_IDS = tuple(t.id for t in TEMPLATES)

# Active compiled rules; replaced in main() after CLI/prompt resolve.
PATTERNS: list[tuple[re.Pattern[str], str]] = [
    (re.compile(t.regex, t.flags), t.replacement)
    for t in TEMPLATES
    if t.id in FINANCE_TEMPLATE_IDS
]


def scrub(value: str) -> tuple[str, int]:
    """Replace sensitive values. Returns (scrubbed_text, replacement_count)."""
    total = 0
    for pattern, replacement in PATTERNS:
        value, n = pattern.subn(replacement, value)
        total += n
    return value, total


def iter_sensitive_matches(text: str) -> list[tuple[str, str]]:
    """Return (matched_text, replacement) pairs found in text."""
    found: list[tuple[str, str]] = []
    for pattern, replacement in PATTERNS:
        for match in pattern.finditer(text):
            found.append((match.group(0), replacement))
    return found


def count_sensitive_replacements(text: str) -> int:
    """Count replacements scrub() would make (same sequential rules)."""
    return scrub(text)[1]


def _require_questionary():
    try:
        import questionary
    except ImportError as exc:
        raise SystemExit(
            "questionary is required for interactive pattern selection. "
            "Install with: pip install questionary"
        ) from exc
    return questionary


def parse_pattern_spec(raw: str) -> tuple[str, str]:
    """Parse 'REGEX=>REPLACEMENT'. Raises ValueError on bad input."""
    if PATTERN_SEP not in raw:
        raise ValueError(
            f"Custom pattern must be REGEX{PATTERN_SEP}REPLACEMENT, got: {raw}"
        )
    regex, replacement = raw.split(PATTERN_SEP, 1)
    regex = regex.strip()
    replacement = replacement.strip()
    if not regex:
        raise ValueError("Custom pattern regex is empty")
    try:
        re.compile(regex)
    except re.error as exc:
        raise ValueError(f"Invalid regex {regex!r}: {exc}") from exc
    return regex, replacement


def compile_pattern_specs(
    specs: list[tuple[str, str, int]],
) -> list[tuple[re.Pattern[str], str]]:
    compiled: list[tuple[re.Pattern[str], str]] = []
    for regex, replacement, flags in specs:
        compiled.append((re.compile(regex, flags), replacement))
    return compiled


def normalize_template_ids(names: list[str]) -> list[str]:
    """Deduplicate template ids; raise ValueError on unknown names."""
    seen: list[str] = []
    unknown: list[str] = []
    for name in names:
        key = name.strip().lower()
        if key not in TEMPLATES_BY_ID:
            unknown.append(name)
            continue
        if key not in seen:
            seen.append(key)
    if unknown:
        valid = ", ".join(TEMPLATE_IDS)
        bad = ", ".join(unknown)
        raise ValueError(f"Unknown template(s): {bad}. Valid: {valid}")
    return seen


def specs_from_templates(template_ids: list[str]) -> list[tuple[str, str, int]]:
    return [
        (TEMPLATES_BY_ID[tid].regex, TEMPLATES_BY_ID[tid].replacement, TEMPLATES_BY_ID[tid].flags)
        for tid in template_ids
    ]


def prompt_template_ids(*, replace: bool) -> list[str]:
    questionary = _require_questionary()
    from questionary.prompts import common as qcommon

    qcommon.INDICATOR_SELECTED = "✅"
    qcommon.INDICATOR_UNSELECTED = "○ "
    checkbox_style = questionary.Style(
        [
            ("selected", "fg:ansigreen bold"),
        ]
    )
    choices = [
        questionary.Choice(
            title=f"{t.id}  {t.description}  → {t.replacement}",
            value=t.id,
            checked=(not replace and t.id in FINANCE_TEMPLATE_IDS),
        )
        for t in TEMPLATES
    ]

    def _enough(selected: list[str]) -> bool | str:
        if replace or selected:
            return True
        return "Select at least one pattern."

    chosen = questionary.checkbox(
        "Which patterns?",
        choices=choices,
        validate=_enough,
        instruction="(space to toggle, enter to confirm)",
        style=checkbox_style,
    ).ask()
    if chosen is None:
        raise typer.Exit(130)
    return chosen


def prompt_custom_specs() -> list[tuple[str, str, int]]:
    extra: list[tuple[str, str, int]] = []
    while True:
        raw = Prompt.ask(
            f"[bold]Add extra pattern[/bold] (regex{PATTERN_SEP}replacement, blank to finish)",
            default="",
        ).strip()
        if not raw:
            break
        try:
            regex, replacement = parse_pattern_spec(raw)
        except ValueError as exc:
            console.print(f"[red]{exc}[/red]")
            continue
        extra.append((regex, replacement, 0))
    return extra


def resolve_pattern_specs(
    cli_templates: list[str] | None,
    cli_patterns: list[str] | None,
    *,
    replace: bool,
) -> tuple[list[str], int]:
    """
    Return (template_ids, custom_count) and set module-level PATTERNS.

    If --template is omitted, prompt with a checkbox list. Custom -p skips the extra-pattern loop.
    """
    try:
        if cli_templates is not None:
            selected = normalize_template_ids(cli_templates)
            if not replace:
                merged: list[str] = []
                for tid in (*FINANCE_TEMPLATE_IDS, *selected):
                    if tid not in merged:
                        merged.append(tid)
                selected = merged
        else:
            selected = prompt_template_ids(replace=replace)
    except ValueError as exc:
        console.print(f"[red]{exc}[/red]")
        raise typer.Exit(1) from exc

    specs = specs_from_templates(selected)

    custom: list[tuple[str, str, int]] = []
    for raw in cli_patterns or []:
        try:
            regex, replacement = parse_pattern_spec(raw)
        except ValueError as exc:
            console.print(f"[red]{exc}[/red]")
            raise typer.Exit(1) from exc
        custom.append((regex, replacement, 0))

    if not cli_patterns:
        custom.extend(prompt_custom_specs())

    if replace and not selected and not custom:
        console.print(
            "[red]No patterns selected. Choose a template or add a custom "
            f"regex{PATTERN_SEP}replacement.[/red]"
        )
        raise typer.Exit(1)

    specs.extend(custom)
    global PATTERNS
    PATTERNS = compile_pattern_specs(specs)
    return selected, len(custom)


def default_output_dir(input_path: Path) -> Path:
    """Place a 'redacted' folder next to a file, or inside an input directory."""
    if input_path.is_file():
        return input_path.parent / "redacted"
    return input_path / "redacted"


def resolve_input_path(raw: str | None) -> Path:
    cwd = str(Path.cwd())
    while True:
        value = raw or Prompt.ask(
            "[bold]Input[/bold] file or directory",
            default=cwd,
        )
        path = Path(value).expanduser().resolve()
        if not path.exists():
            console.print(f"[red]Path does not exist:[/red] {path}")
        elif path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            return path
        elif path.is_dir():
            return path
        else:
            kinds = ", ".join(sorted(s.lstrip(".").upper() for s in SUPPORTED_SUFFIXES))
            console.print(
                f"[red]Not a supported file ({kinds}) or directory:[/red] {path}",
            )
        raw = None


def resolve_output_dir(raw: str | None, input_path: Path) -> Path:
    default = default_output_dir(input_path)
    value = raw or Prompt.ask(
        "[bold]Output[/bold] directory",
        default=str(default),
    )
    return Path(value).expanduser().resolve()


def has_subdirectories(path: Path) -> bool:
    return path.is_dir() and any(p.is_dir() for p in path.iterdir())


def resolve_recursive(recursive: bool | None, input_path: Path) -> bool:
    """Return whether to scan subdirs; prompt when present and flag was omitted."""
    if input_path.is_file():
        return False
    if recursive is not None:
        return recursive
    if not has_subdirectories(input_path):
        return False
    return Confirm.ask(
        "[bold]Subdirectories found.[/bold] Scan recursively?",
        default=False,
    )


def _is_under(path: Path, root: Path | None) -> bool:
    if root is None:
        return False
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def collect_files(
    input_path: Path,
    *,
    recursive: bool = False,
    skip_under: Path | None = None,
) -> list[Path]:
    if input_path.is_file():
        return [input_path]
    iterator = input_path.rglob("*") if recursive else input_path.iterdir()
    files = [
        p
        for p in iterator
        if p.is_file()
        and p.suffix.lower() in SUPPORTED_SUFFIXES
        and not _is_under(p, skip_under)
    ]
    return sorted(files, key=lambda p: (p.suffix.lower(), str(p).lower()))


def source_label(src: Path, input_path: Path) -> str:
    """Path shown in the results table (relative when scanning a directory)."""
    if input_path.is_file():
        return src.name
    try:
        return str(src.relative_to(input_path))
    except ValueError:
        return src.name


def destination_for(
    src: Path,
    input_path: Path,
    output_dir: Path,
    *,
    recursive: bool,
) -> Path:
    """Mirror subdirectory layout under output_dir when scanning recursively."""
    if input_path.is_file() or not recursive:
        return output_dir / src.name
    return output_dir / src.relative_to(input_path)


def confirm_write(dest: Path, *, overwrite: bool) -> bool:
    """Return True if dest may be written. Existing files require confirmation."""
    if not dest.exists() or overwrite:
        return True
    return Confirm.ask(
        f"[yellow]Overwrite existing file?[/yellow] {dest}",
        default=False,
    )


def read_csv_rows(src: Path, encoding: str) -> tuple[list[dict[str, str]], list[str]]:
    """Return (rows, kept_columns)."""
    with src.open(newline="", encoding=encoding) as fin:
        sample = fin.read(2048)
        fin.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        rows = list(csv.DictReader(fin, dialect=dialect))
    if not rows:
        return [], []
    cols = [c for c in rows[0] if c.strip().lower() in KEEP]
    return rows, cols


def count_csv_redactions(src: Path, encoding: str) -> tuple[int, int, int]:
    """Return (row_count, column_count, redaction_count) without writing."""
    rows, cols = read_csv_rows(src, encoding)
    if not rows:
        return 0, 0, 0
    redactions = 0
    for row in rows:
        for col in cols:
            redactions += count_sensitive_replacements(row[col])
    return len(rows), len(cols), redactions


def redact_csv(src: Path, dest: Path, encoding: str) -> tuple[int, int, int]:
    """Write a redacted CSV. Returns (row_count, column_count, redaction_count)."""
    rows, cols = read_csv_rows(src, encoding)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        dest.write_text("", encoding=encoding)
        return 0, 0, 0

    redactions = 0
    with dest.open("w", newline="", encoding=encoding) as fout:
        writer = csv.writer(fout)
        writer.writerow(cols)
        for row in rows:
            out_row = []
            for col in cols:
                scrubbed, n = scrub(row[col])
                redactions += n
                out_row.append(scrubbed)
            writer.writerow(out_row)
    return len(rows), len(cols), redactions


def line_count(text: str) -> int:
    if not text:
        return 0
    return text.count("\n") + (0 if text.endswith("\n") else 1)


def write_encoding(encoding: str) -> str:
    """ElementTree/JSON prefer plain utf-8 over utf-8-sig for the declared encoding."""
    return "utf-8" if encoding.lower().replace("_", "-") == "utf-8-sig" else encoding


def count_text_redactions(src: Path, encoding: str) -> tuple[int, int]:
    """Return (line_count, redaction_count) without writing."""
    text = src.read_text(encoding=encoding)
    return line_count(text), count_sensitive_replacements(text)


def redact_text(src: Path, dest: Path, encoding: str) -> tuple[int, int]:
    """Write a redacted text file. Returns (line_count, redaction_count)."""
    text = src.read_text(encoding=encoding)
    scrubbed, redactions = scrub(text)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(scrubbed, encoding=encoding)
    return line_count(text), redactions


def scrub_scalar(value: Any) -> tuple[Any, int]:
    """Scrub strings; convert sensitive integers to placeholder strings."""
    if isinstance(value, str):
        return scrub(value)
    if isinstance(value, bool) or value is None:
        return value, 0
    if isinstance(value, int):
        scrubbed, n = scrub(str(value))
        return (scrubbed, n) if n else (value, 0)
    if isinstance(value, float) and value.is_integer():
        return scrub_scalar(int(value))
    return value, 0


def scrub_json_value(value: Any) -> tuple[Any, int]:
    """Recursively scrub JSON-compatible values. Returns (value, redaction_count)."""
    if isinstance(value, dict):
        total = 0
        out: dict[Any, Any] = {}
        for key, item in value.items():
            new_key, key_n = scrub(key) if isinstance(key, str) else (key, 0)
            new_item, item_n = scrub_json_value(item)
            out[new_key] = new_item
            total += key_n + item_n
        return out, total
    if isinstance(value, list):
        total = 0
        out_list: list[Any] = []
        for item in value:
            new_item, item_n = scrub_json_value(item)
            out_list.append(new_item)
            total += item_n
        return out_list, total
    return scrub_scalar(value)


def count_json_redactions(src: Path, encoding: str) -> tuple[int, int]:
    """Return (top_level_item_count, redaction_count) without writing."""
    data = json.loads(src.read_text(encoding=encoding))
    _, redactions = scrub_json_value(data)
    if isinstance(data, list):
        return len(data), redactions
    if isinstance(data, dict):
        return len(data), redactions
    return 1, redactions


def redact_json(src: Path, dest: Path, encoding: str) -> tuple[int, int]:
    """Write a redacted JSON file. Returns (top_level_item_count, redaction_count)."""
    data = json.loads(src.read_text(encoding=encoding))
    scrubbed, redactions = scrub_json_value(data)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(
        json.dumps(scrubbed, ensure_ascii=False, indent=2) + "\n",
        encoding=write_encoding(encoding),
    )
    if isinstance(data, list):
        return len(data), redactions
    if isinstance(data, dict):
        return len(data), redactions
    return 1, redactions


def scrub_xml_element(elem: ET.Element) -> int:
    """Scrub text/tails/attributes on an element tree in-place."""
    total = 0
    if elem.text:
        elem.text, n = scrub(elem.text)
        total += n
    if elem.tail:
        elem.tail, n = scrub(elem.tail)
        total += n
    for key, val in list(elem.attrib.items()):
        scrubbed, n = scrub(val)
        elem.attrib[key] = scrubbed
        total += n
    for child in elem:
        total += scrub_xml_element(child)
    return total


def _xml_detail(root: ET.Element) -> str:
    return f"{sum(1 for _ in root.iter())} nodes"


def count_xml_redactions(src: Path, encoding: str) -> tuple[str, int]:
    """Return (detail, redaction_count) without writing."""
    root = ET.parse(src).getroot()
    return _xml_detail(root), scrub_xml_element(root)


def redact_xml(src: Path, dest: Path, encoding: str) -> tuple[str, int]:
    """Write a redacted XML file. Returns (detail, redaction_count)."""
    tree = ET.parse(src)
    root = tree.getroot()
    redactions = scrub_xml_element(root)
    dest.parent.mkdir(parents=True, exist_ok=True)
    tree.write(dest, encoding=write_encoding(encoding), xml_declaration=True)
    return _xml_detail(root), redactions


def _require_bs4():
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise SystemExit(
            "BeautifulSoup is required for HTML redaction. "
            "Install with: pip install beautifulsoup4"
        ) from exc
    return BeautifulSoup


def scrub_html_soup(soup: Any) -> int:
    """Scrub text nodes and string attributes in a BeautifulSoup tree."""
    total = 0
    for node in list(soup.find_all(string=True)):
        text = str(node)
        scrubbed, n = scrub(text)
        if n:
            node.replace_with(scrubbed)
            total += n
    for tag in soup.find_all(True):
        for key, val in list(tag.attrs.items()):
            if isinstance(val, str):
                scrubbed, n = scrub(val)
                if n:
                    tag.attrs[key] = scrubbed
                    total += n
            elif isinstance(val, list):
                new_list: list[str] = []
                changed = False
                for item in val:
                    if isinstance(item, str):
                        scrubbed, n = scrub(item)
                        new_list.append(scrubbed)
                        if n:
                            total += n
                            changed = True
                    else:
                        new_list.append(item)
                if changed:
                    tag.attrs[key] = new_list
    return total


def count_html_redactions(src: Path, encoding: str) -> tuple[str, int]:
    """Return (detail, redaction_count) without writing."""
    BeautifulSoup = _require_bs4()
    text = src.read_text(encoding=encoding)
    soup = BeautifulSoup(text, "html.parser")
    redactions = scrub_html_soup(soup)
    return f"{line_count(text)} lines", redactions


def redact_html(src: Path, dest: Path, encoding: str) -> tuple[str, int]:
    """Write a redacted HTML file. Returns (detail, redaction_count)."""
    BeautifulSoup = _require_bs4()
    text = src.read_text(encoding=encoding)
    soup = BeautifulSoup(text, "html.parser")
    redactions = scrub_html_soup(soup)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(str(soup), encoding=encoding)
    return f"{line_count(text)} lines", redactions


def _require_openpyxl():
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit(
            "openpyxl is required for XLSX redaction. Install with: pip install openpyxl"
        ) from exc
    return load_workbook


def scrub_xlsx_workbook(wb: Any) -> tuple[int, int, int]:
    """
    Scrub all cell values in an openpyxl workbook in-place.

    Returns (sheet_count, cell_count, redaction_count).
    """
    cells = 0
    redactions = 0
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cells += 1
                scrubbed, n = scrub_scalar(cell.value)
                if n:
                    cell.value = scrubbed
                    redactions += n
    return len(wb.worksheets), cells, redactions


def count_xlsx_redactions(src: Path) -> tuple[int, int, int]:
    """Return (sheet_count, cell_count, redaction_count) without writing."""
    load_workbook = _require_openpyxl()
    wb = load_workbook(src)
    try:
        return scrub_xlsx_workbook(wb)
    finally:
        wb.close()


def redact_xlsx(src: Path, dest: Path) -> tuple[int, int, int]:
    """Write a redacted XLSX file. Returns (sheet_count, cell_count, redaction_count)."""
    load_workbook = _require_openpyxl()
    wb = load_workbook(src)
    try:
        sheets, cells, redactions = scrub_xlsx_workbook(wb)
        dest.parent.mkdir(parents=True, exist_ok=True)
        wb.save(dest)
        return sheets, cells, redactions
    finally:
        wb.close()


class UnredactableError(Exception):
    """Source has no usable text (and OCR could not recover it)."""

    def __init__(self, reason: str):
        super().__init__(reason)
        self.reason = reason


@dataclass(frozen=True)
class OcrConfig:
    enabled: bool = True
    tessdata: str | None = None
    language: str = DEFAULT_OCR_LANG


@dataclass
class PdfTextReport:
    page_count: int
    usable_pages: int
    empty_text_pages: list[int]  # 1-based page numbers still without text
    match_count: int
    image_pages: int
    ocr_pages: int = 0

    @property
    def has_usable_text(self) -> bool:
        return self.usable_pages > 0

    @property
    def is_likely_scanned(self) -> bool:
        return not self.has_usable_text and self.image_pages > 0


def _require_pymupdf():
    try:
        import pymupdf
    except ImportError as exc:
        raise SystemExit(
            "PyMuPDF is required for PDF and image redaction. "
            "Install with: pip install pymupdf"
        ) from exc
    return pymupdf


def _page_alnum_count(text: str) -> int:
    return sum(1 for c in text if c.isalnum())


def _ocr_missing_reason(*, enabled: bool) -> str:
    if not enabled:
        return (
            "no extractable text layer (likely scanned/image-only) — "
            "OCR is off (--no-ocr); sensitive data may still be visible in images"
        )
    return (
        "no extractable text layer (likely scanned/image-only) — "
        "OCR needs Tesseract language data (tessdata). "
        "Install Tesseract, set TESSDATA_PREFIX, or pass --tessdata"
    )


def resolve_ocr_config(ocr: OcrConfig) -> OcrConfig:
    """Fill tessdata path and drop language codes that are not installed."""
    if not ocr.enabled:
        return ocr
    pymupdf = _require_pymupdf()
    try:
        tessdata = pymupdf.get_tessdata(ocr.tessdata)
    except Exception:
        return OcrConfig(enabled=True, tessdata=None, language=ocr.language)
    return OcrConfig(
        enabled=True,
        tessdata=tessdata,
        language=_ocr_languages_in(tessdata, ocr.language),
    )


def _ocr_languages_in(tessdata: str, requested: str) -> str:
    root = Path(tessdata)
    kept = [
        code
        for code in requested.split("+")
        if code and (root / f"{code}.traineddata").is_file()
    ]
    if kept:
        return "+".join(kept)
    if (root / "eng.traineddata").is_file():
        return "eng"
    return requested


def _pixmap_for_ocr(pymupdf: Any, pix: Any) -> Any:
    if pix.alpha:
        return pymupdf.Pixmap(pix, 0)
    return pix


def ocr_textpage(page: Any, pymupdf: Any, ocr: OcrConfig) -> Any:
    if not ocr.tessdata:
        raise UnredactableError(_ocr_missing_reason(enabled=ocr.enabled))
    return page.get_textpage_ocr(
        dpi=OCR_DPI,
        full=True,
        language=ocr.language,
        tessdata=ocr.tessdata,
    )


def _add_match_redactions(
    page: Any,
    pymupdf: Any,
    hits: list[tuple[str, str]],
    *,
    textpage: Any | None = None,
) -> int:
    count = 0
    seen: set[str] = set()
    search_kw: dict[str, Any] = {}
    if textpage is not None:
        search_kw["textpage"] = textpage
    for matched_text, replacement in hits:
        if matched_text in seen:
            continue
        seen.add(matched_text)
        for rect in page.search_for(matched_text, **search_kw):
            page.add_redact_annot(
                rect,
                text=replacement,
                fill=(0, 0, 0),
                text_color=(1, 1, 1),
                align=pymupdf.TEXT_ALIGN_CENTER,
                cross_out=False,
            )
            count += 1
    return count


def _apply_page_redactions(page: Any, pymupdf: Any, *, covering_images: bool) -> None:
    page.apply_redactions(
        images=(
            pymupdf.PDF_REDACT_IMAGE_PIXELS
            if covering_images
            else pymupdf.PDF_REDACT_IMAGE_NONE
        ),
        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
    )


def _unredactable_reason(report: PdfTextReport, ocr: OcrConfig) -> str:
    if report.is_likely_scanned:
        return _ocr_missing_reason(enabled=ocr.enabled)
    return "no extractable text layer — cannot locate or remove sensitive values"


def analyze_pdf(src: Path, ocr: OcrConfig | None = None) -> PdfTextReport:
    """Inspect extractable text / images without modifying the file."""
    ocr = ocr or OcrConfig(enabled=False)
    pymupdf = _require_pymupdf()
    doc = pymupdf.open(src)
    try:
        empty_text_pages: list[int] = []
        usable_pages = 0
        match_count = 0
        image_pages = 0
        ocr_pages = 0
        for index, page in enumerate(doc, start=1):
            text = page.get_text()
            if page.get_images(full=True):
                image_pages += 1
            if _page_alnum_count(text) >= MIN_PAGE_ALNUM:
                usable_pages += 1
                match_count += len(iter_sensitive_matches(text))
                continue
            if ocr.enabled and ocr.tessdata:
                try:
                    textpage = ocr_textpage(page, pymupdf, ocr)
                    text = page.get_text(textpage=textpage)
                except Exception:
                    empty_text_pages.append(index)
                    continue
                ocr_pages += 1
                if _page_alnum_count(text) >= MIN_PAGE_ALNUM:
                    usable_pages += 1
                    match_count += len(iter_sensitive_matches(text))
                else:
                    empty_text_pages.append(index)
                continue
            empty_text_pages.append(index)
        return PdfTextReport(
            page_count=doc.page_count,
            usable_pages=usable_pages,
            empty_text_pages=empty_text_pages,
            match_count=match_count,
            image_pages=image_pages,
            ocr_pages=ocr_pages,
        )
    finally:
        doc.close()


def redact_pdf(
    src: Path,
    dest: Path,
    ocr: OcrConfig,
) -> tuple[int, int, list[int], int]:
    """
    Physically remove matched text via PyMuPDF redaction annotations.

    Returns (page_count, redaction_count, empty_text_pages, ocr_pages).
    Raises UnredactableError when the PDF has no usable text layer and OCR cannot help.
    """
    pymupdf = _require_pymupdf()
    # Tighter glyph boxes reduce accidental removal of neighboring lines.
    pymupdf.TOOLS.set_small_glyph_heights(True)

    doc = pymupdf.open(src)
    redaction_count = 0
    empty_text_pages: list[int] = []
    ocr_pages = 0
    usable_pages = 0
    image_pages = 0
    try:
        for index, page in enumerate(doc, start=1):
            if page.get_images(full=True):
                image_pages += 1
            native = page.get_text()
            textpage = None
            covering_images = False
            text = native
            if _page_alnum_count(native) < MIN_PAGE_ALNUM:
                if ocr.enabled and ocr.tessdata:
                    try:
                        textpage = ocr_textpage(page, pymupdf, ocr)
                        text = page.get_text(textpage=textpage)
                        covering_images = True
                        ocr_pages += 1
                    except Exception:
                        empty_text_pages.append(index)
                        continue
                else:
                    empty_text_pages.append(index)
                    continue
                if _page_alnum_count(text) < MIN_PAGE_ALNUM:
                    empty_text_pages.append(index)
                else:
                    usable_pages += 1
            else:
                usable_pages += 1
            hits = iter_sensitive_matches(text)
            redaction_count += _add_match_redactions(
                page, pymupdf, hits, textpage=textpage
            )
            # Remove overlapping text from the content stream (not just cover it).
            _apply_page_redactions(page, pymupdf, covering_images=covering_images)

        if usable_pages == 0:
            report = PdfTextReport(
                page_count=doc.page_count,
                usable_pages=0,
                empty_text_pages=empty_text_pages,
                match_count=0,
                image_pages=image_pages,
                ocr_pages=ocr_pages,
            )
            raise UnredactableError(_unredactable_reason(report, ocr))

        dest.parent.mkdir(parents=True, exist_ok=True)
        # Strip Title/Author/Subject/etc. and XMP so they cannot leak identifiers.
        doc.set_metadata({})
        doc.del_xml_metadata()
        # garbage/deflate purge removed content so it is not extractable later.
        doc.save(dest, garbage=4, deflate=True, clean=True)
        return doc.page_count, redaction_count, empty_text_pages, ocr_pages
    finally:
        doc.close()


def redact_image(
    src: Path,
    dest: Path,
    ocr: OcrConfig,
    *,
    write: bool,
) -> tuple[str, int]:
    """OCR an image, blank matched pixels, optionally write dest. Returns (detail, count)."""
    if not ocr.enabled or not ocr.tessdata:
        if not ocr.enabled:
            raise UnredactableError("OCR is off (--no-ocr); cannot redact image pixels")
        raise UnredactableError(
            "image OCR needs Tesseract language data (tessdata). "
            "Install Tesseract, set TESSDATA_PREFIX, or pass --tessdata"
        )

    pymupdf = _require_pymupdf()
    pymupdf.TOOLS.set_small_glyph_heights(True)
    pix = _pixmap_for_ocr(pymupdf, pymupdf.Pixmap(src))
    width, height = pix.width, pix.height
    try:
        ocr_pdf = pymupdf.open(
            "pdf",
            pix.pdfocr_tobytes(language=ocr.language, tessdata=ocr.tessdata),
        )
    finally:
        pix = None

    try:
        page = ocr_pdf[0]
        hits = iter_sensitive_matches(page.get_text())
        redactions = _add_match_redactions(page, pymupdf, hits)
        if write:
            _apply_page_redactions(page, pymupdf, covering_images=True)
            dest.parent.mkdir(parents=True, exist_ok=True)
            out = page.get_pixmap(
                matrix=pymupdf.Matrix(
                    width / page.rect.width,
                    height / page.rect.height,
                ),
                alpha=False,
            )
            out.save(dest)
        return f"{width}×{height}", redactions
    finally:
        ocr_pdf.close()


@dataclass
class FileWork:
    detail: str
    redactions: int
    empty_pages: list[int] = field(default_factory=list)
    skip_reason: str | None = None
    used_ocr: bool = False


def format_file_counts(files: list[Path]) -> str:
    counts = Counter(p.suffix.lower().lstrip(".").upper() for p in files)
    parts: list[str] = []
    for label in _FILE_COUNT_ORDER:
        n = counts.pop(label, 0)
        if not n:
            continue
        style = _FILE_COUNT_STYLE.get(label, "white")
        parts.append(f"[{style}]{n} {label}[/{style}]")
    for label, n in sorted(counts.items()):
        parts.append(f"{n} {label}")
    return ", ".join(parts)


def process_source(
    src: Path,
    dest: Path,
    encoding: str,
    *,
    write: bool,
    ocr: OcrConfig,
) -> FileWork:
    """Inspect or redact one file. Unredactable sources return skip_reason instead of raising."""
    suffix = src.suffix.lower()
    if suffix in PDF_SUFFIXES:
        if not write:
            report = analyze_pdf(src, ocr)
            work = FileWork(
                detail=f"{report.page_count}p",
                redactions=report.match_count,
                empty_pages=report.empty_text_pages,
                used_ocr=report.ocr_pages > 0,
            )
            if not report.has_usable_text:
                work.skip_reason = _unredactable_reason(report, ocr)
            return work
        try:
            pages, redactions, empty_pages, ocr_pages = redact_pdf(src, dest, ocr)
        except UnredactableError as exc:
            return FileWork(detail="—", redactions=0, skip_reason=exc.reason)
        return FileWork(
            detail=f"{pages}p",
            redactions=redactions,
            empty_pages=empty_pages,
            used_ocr=ocr_pages > 0,
        )
    if suffix in IMAGE_SUFFIXES:
        try:
            detail, redactions = redact_image(src, dest, ocr, write=write)
        except UnredactableError as exc:
            return FileWork(detail="—", redactions=0, skip_reason=exc.reason)
        return FileWork(detail=detail, redactions=redactions, used_ocr=True)
    if suffix in CSV_SUFFIXES:
        if write:
            rows, cols, redactions = redact_csv(src, dest, encoding)
        else:
            rows, cols, redactions = count_csv_redactions(src, encoding)
        return FileWork(
            detail=f"{rows} rows / {cols} cols",
            redactions=redactions,
        )
    if suffix in TEXT_SUFFIXES:
        if write:
            lines, redactions = redact_text(src, dest, encoding)
        else:
            lines, redactions = count_text_redactions(src, encoding)
        return FileWork(detail=f"{lines} lines", redactions=redactions)
    if suffix in JSON_SUFFIXES:
        if write:
            items, redactions = redact_json(src, dest, encoding)
        else:
            items, redactions = count_json_redactions(src, encoding)
        return FileWork(detail=f"{items} items", redactions=redactions)
    if suffix in XML_SUFFIXES:
        if write:
            detail, redactions = redact_xml(src, dest, encoding)
        else:
            detail, redactions = count_xml_redactions(src, encoding)
        return FileWork(detail=detail, redactions=redactions)
    if suffix in HTML_SUFFIXES:
        if write:
            detail, redactions = redact_html(src, dest, encoding)
        else:
            detail, redactions = count_html_redactions(src, encoding)
        return FileWork(detail=detail, redactions=redactions)
    if suffix in XLSX_SUFFIXES:
        if write:
            sheets, cells, redactions = redact_xlsx(src, dest)
        else:
            sheets, cells, redactions = count_xlsx_redactions(src)
        return FileWork(
            detail=f"{sheets} sheets / {cells} cells",
            redactions=redactions,
        )
    raise ValueError(f"Unsupported file type: {suffix}")


@app.command()
def main(
    input: Annotated[
        Path | None,
        typer.Option(
            "--input",
            "-i",
            help=(
                "Supported file or directory "
                "(CSV/TSV/TXT/JSON/XML/HTML/XLSX/PDF/PNG/JPEG/TIFF, OFX/QIF). "
                "Prompt default: current directory."
            ),
            show_default=False,
        ),
    ] = None,
    output: Annotated[
        Path | None,
        typer.Option(
            "--output",
            "-o",
            help=(
                "Output directory (prompt default: redacted next to a file, "
                "or inside an input directory)."
            ),
            show_default=False,
        ),
    ] = None,
    encoding: Annotated[
        str,
        typer.Option(
            "--encoding",
            "-e",
            help="Text encoding for text-based formats (not PDF/XLSX/images).",
        ),
    ] = "utf-8-sig",
    dry_run: Annotated[
        bool,
        typer.Option(
            "--dry-run",
            "-n",
            help="List files that would be written without creating them.",
        ),
    ] = False,
    overwrite: Annotated[
        bool,
        typer.Option(
            "--overwrite",
            "-y",
            help="Overwrite existing output files without asking.",
        ),
    ] = False,
    recursive: Annotated[
        bool | None,
        typer.Option(
            "--recursive/--no-recursive",
            "-r",
            help=(
                "Include files in subdirectories. "
                "Prompted when omitted and subdirectories are present."
            ),
            show_default=False,
        ),
    ] = None,
    template: Annotated[
        list[str] | None,
        typer.Option(
            "--template",
            "-t",
            help=(
                "Named pattern template to enable "
                f"({', '.join(TEMPLATE_IDS)}). Repeatable. "
                "Finance templates stay on unless --replace-patterns. "
                "Prompted when omitted."
            ),
            show_default=False,
        ),
    ] = None,
    pattern: Annotated[
        list[str] | None,
        typer.Option(
            "--pattern",
            "-p",
            help=(
                f"Custom rule as REGEX{PATTERN_SEP}REPLACEMENT. Repeatable. "
                "Prompted when omitted."
            ),
            show_default=False,
        ),
    ] = None,
    replace_patterns: Annotated[
        bool,
        typer.Option(
            "--replace-patterns",
            help="Do not auto-include the default finance templates.",
        ),
    ] = False,
    ocr: Annotated[
        bool,
        typer.Option(
            "--ocr/--no-ocr",
            help=(
                "OCR image files and scanned PDF pages. "
                "Needs Tesseract language data (tessdata)."
            ),
        ),
    ] = True,
    tessdata: Annotated[
        Path | None,
        typer.Option(
            "--tessdata",
            help="Tesseract tessdata folder (default: TESSDATA_PREFIX or Tesseract install).",
            show_default=False,
        ),
    ] = None,
    ocr_lang: Annotated[
        str,
        typer.Option(
            "--ocr-lang",
            help="Tesseract language codes, '+' separated (missing codes are dropped).",
        ),
    ] = DEFAULT_OCR_LANG,
) -> None:
    """Redact sensitive fields from Scandinavian financial exports."""
    input_path = resolve_input_path(str(input) if input is not None else None)
    output_dir = resolve_output_dir(
        str(output) if output is not None else None,
        input_path,
    )
    recursive = resolve_recursive(recursive, input_path)
    selected_templates, custom_count = resolve_pattern_specs(
        template,
        pattern,
        replace=replace_patterns,
    )
    skip_under = output_dir if _is_under(output_dir, input_path) else None
    files = collect_files(input_path, recursive=recursive, skip_under=skip_under)

    if not files:
        kinds = ", ".join(sorted(s.lstrip(".").upper() for s in SUPPORTED_SUFFIXES))
        scope = "recursively under" if recursive else "in"
        console.print(f"[yellow]No {kinds} files found {scope}[/yellow] {input_path}")
        raise typer.Exit(1)

    if tessdata is not None and not tessdata.is_dir():
        console.print(f"[red]tessdata folder not found:[/red] {tessdata}")
        raise typer.Exit(1)

    ocr_config = OcrConfig(
        enabled=ocr,
        tessdata=str(tessdata) if tessdata is not None else None,
        language=ocr_lang,
    )
    needs_ocr = any(p.suffix.lower() in IMAGE_SUFFIXES | PDF_SUFFIXES for p in files)
    if needs_ocr and ocr_config.enabled:
        ocr_config = resolve_ocr_config(ocr_config)

    counts = format_file_counts(files)
    mode_bits = []
    if recursive:
        mode_bits.append("[dim](recursive)[/dim]")
    if dry_run:
        mode_bits.append("[dim](dry-run)[/dim]")
    if replace_patterns:
        mode_bits.append("[dim](replace patterns)[/dim]")
    if needs_ocr and not ocr_config.enabled:
        mode_bits.append("[yellow](OCR off)[/yellow]")
    elif needs_ocr and ocr_config.tessdata:
        mode_bits.append(f"[dim](OCR {ocr_config.language})[/dim]")
    elif needs_ocr:
        mode_bits.append("[yellow](OCR unavailable — install Tesseract tessdata)[/yellow]")
    template_note = ", ".join(selected_templates) if selected_templates else "(none)"
    if custom_count:
        template_note += f" + {custom_count} custom"
    console.print(
        Panel.fit(
            f"[bold]Input[/bold]  {input_path}\n"
            f"[bold]Output[/bold] {output_dir}\n"
            f"[bold]Files[/bold]  {len(files)} ({counts})\n"
            f"[bold]Patterns[/bold] {template_note}"
            + (("  " + " ".join(mode_bits)) if mode_bits else ""),
            title="pyredact",
            border_style="cyan",
        )
    )

    results = Table(show_header=True, header_style="bold")
    results.add_column("Type")
    results.add_column("Source")
    results.add_column("Destination")
    results.add_column("Detail", justify="right")
    results.add_column("Redactions", justify="right")
    results.add_column("Status")

    not_redacted: list[tuple[str, str]] = []
    partial_warnings: list[tuple[str, list[int]]] = []
    wrote_any = False
    skipped_overwrite = 0
    total_redactions = 0
    files_with_redactions = 0

    with Progress(
        SpinnerColumn(),
        TextColumn("[progress.description]{task.description}"),
        BarColumn(),
        TaskProgressColumn(),
        console=console,
        transient=True,
    ) as progress:
        task = progress.add_task("Redacting", total=len(files))
        for src in files:
            dest = destination_for(
                src, input_path, output_dir, recursive=recursive
            )
            label = source_label(src, input_path)
            kind = src.suffix.lower().lstrip(".").upper()
            try:
                if not dry_run and not confirm_write(dest, overwrite=overwrite):
                    skipped_overwrite += 1
                    results.add_row(
                        kind,
                        label,
                        str(dest),
                        "—",
                        "—",
                        "[yellow]skipped (exists)[/yellow]",
                    )
                    continue

                work = process_source(
                    src, dest, encoding, write=not dry_run, ocr=ocr_config
                )
                if work.skip_reason:
                    not_redacted.append((label, work.skip_reason))
                    results.add_row(
                        kind,
                        label,
                        "—",
                        work.detail,
                        str(work.redactions),
                        "[bold red]NOT REDACTED[/bold red]",
                    )
                    continue

                if not dry_run:
                    wrote_any = True
                total_redactions += work.redactions
                if work.redactions:
                    files_with_redactions += 1

                exists_note = (
                    " [yellow](would overwrite)[/yellow]"
                    if dry_run and dest.exists()
                    else ""
                )
                status = (
                    f"[cyan]would write[/cyan]{exists_note}"
                    if dry_run
                    else "[green]ok[/green]"
                )
                if work.used_ocr:
                    status += " [dim](OCR)[/dim]"
                if work.empty_pages:
                    partial_warnings.append((label, work.empty_pages))
                    status += " [yellow](partial text)[/yellow]"
                results.add_row(
                    kind,
                    label,
                    str(dest),
                    work.detail,
                    str(work.redactions),
                    status,
                )
            except SystemExit:
                raise
            except Exception as exc:  # noqa: BLE001 - surface per-file errors
                results.add_row(
                    kind,
                    label,
                    str(dest),
                    "—",
                    "—",
                    f"[red]error: {exc}[/red]",
                )
            finally:
                progress.advance(task)

    console.print(results)

    summary_bits = [
        f"[bold]{total_redactions}[/bold] redaction(s)",
        f"across [bold]{files_with_redactions}[/bold] file(s)",
    ]
    if skipped_overwrite:
        summary_bits.append(f"[yellow]{skipped_overwrite} skipped (exists)[/yellow]")
    console.print(
        Panel.fit(
            " · ".join(summary_bits),
            title="Summary",
            border_style="cyan",
        )
    )

    if partial_warnings:
        lines = "\n".join(
            f" • [bold]{name}[/bold] — no text on page(s) "
            f"{', '.join(str(p) for p in pages)}"
            for name, pages in partial_warnings
        )
        console.print(
            Panel(
                "[bold yellow]PARTIAL TEXT LAYER[/bold yellow]\n\n"
                "These files were processed, but some pages have little/no extractable "
                "text even after OCR. Anything still only present as unread pixels "
                "was [bold]not[/bold] redacted.\n\n"
                f"{lines}",
                title="Warning",
                border_style="yellow",
            )
        )

    if not_redacted:
        lines = "\n".join(
            f" • [bold]{name}[/bold] — {reason}" for name, reason in not_redacted
        )
        console.print(
            Panel(
                "[bold white on red] NOT REDACTED [/bold white on red]\n\n"
                "These files were [bold]skipped[/bold] and [bold]not[/bold] written to "
                "the output folder. Treat the originals as still sensitive.\n\n"
                f"{lines}\n\n"
                "[dim]Tip: install Tesseract language data (tessdata), or export a "
                "text-based PDF from your bank, then re-run.[/dim]",
                title="⚠ DO NOT SHARE AS REDACTED",
                border_style="red",
            )
        )

    if dry_run:
        console.print("[dim]Dry-run complete — no files written.[/dim]")
    elif wrote_any:
        console.print(f"[green]Done.[/green] Wrote to {output_dir}")
    elif not_redacted and not skipped_overwrite:
        console.print("[red]Nothing was written — no file could be safely redacted.[/red]")
    elif skipped_overwrite and not wrote_any:
        console.print("[yellow]Nothing was written — existing outputs were left untouched.[/yellow]")

    if not_redacted:
        raise typer.Exit(2)


if __name__ == "__main__":
    try:
        app()
    except KeyboardInterrupt:
        console.print("\n[yellow]Interrupted — exiting.[/yellow]")
        raise SystemExit(130) from None
